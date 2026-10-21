#!/usr/bin/env python3
"""
compile_hymns.py

Merges finished rows from Hymns_To_Add.xlsx into hymns_data.json, keeping the
JSON sorted alphabetically by title. Pure local script - no AI calls, no
internet access, safe to run as often as you like.

Usage:
    python3 compile_hymns.py
    python3 compile_hymns.py --dry-run      # preview without writing anything
    python3 compile_hymns.py --xlsx path/to/file.xlsx --json path/to/data.json

What it does each run:
  1. Reads every row in the "To Add" sheet.
  2. A row is "ready" once it has a title AND at least one verse filled in.
     Rows missing verses are left alone (still pending) so you can fill them
     in over time.
  3. Ready rows are converted into hymn records matching the existing JSON
     schema, checked against the current data for likely duplicate titles
     (fuzzy match, just a warning - it will not block you).
  4. New records get the next available sequential id (max existing id + 1,
     incrementing).
  5. The full hymn list (old + new) is re-sorted alphabetically by title and
     written back to hymns_data.json.
  6. Rows that were successfully merged are deleted from "To Add" and logged
     in the "Already Added" sheet, so the spreadsheet always shows only what
     still needs work.
"""

import argparse
import difflib
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

HEADERS = ["title", "tune", "verses", "Author", "Composer", "midi_url", "mp3_url", "credits"]
REQUIRED_FOR_MERGE = ["title", "verses"]  # a row needs at least these to be merged
EXAMPLE_MARKER = "EXAMPLE"


def load_json(json_path: Path) -> list[dict]:
    """Load hymns_data.json, tolerating the one-off trailing-comma glitch
    that can show up if the file was hand-edited."""
    raw = json_path.read_text(encoding="utf-8")
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        fixed = re.sub(r",(\s*\r?\n\s*[\]\}])", r"\1", raw)
        return json.loads(fixed)


def split_verses(cell_text: str) -> list[str]:
    """Turn one spreadsheet cell into the list-of-verse-strings the JSON expects.
    Verses are separated by a blank line; line breaks inside a verse are kept."""
    if not cell_text:
        return []
    normalized = cell_text.replace("\r\n", "\n").strip()
    parts = re.split(r"\n\s*\n", normalized)
    return [p.strip() for p in parts if p.strip()]


def row_to_dict(row_values: tuple) -> dict:
    return {h: (v.strip() if isinstance(v, str) else v) for h, v in zip(HEADERS, row_values)}


def is_example_row(record: dict) -> bool:
    title = record.get("title") or ""
    return title.strip().upper().startswith(EXAMPLE_MARKER)


def is_ready(record: dict) -> bool:
    return all(record.get(f) for f in REQUIRED_FOR_MERGE)


def find_near_duplicates(title: str, existing_titles: list[str], cutoff: float = 0.82) -> list[str]:
    return difflib.get_close_matches(title.lower().strip(), existing_titles, n=2, cutoff=cutoff)


def main():
    parser = argparse.ArgumentParser(description="Merge new hymns from spreadsheet into hymns_data.json")
    parser.add_argument("--xlsx", default="Hymns_To_Add.xlsx", help="Path to the intake spreadsheet")
    parser.add_argument("--json", default="hymns_data.json", help="Path to hymns_data.json")
    parser.add_argument("--dry-run", action="store_true", help="Show what would happen, write nothing")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    json_path = Path(args.json)

    if not xlsx_path.exists():
        sys.exit(f"Spreadsheet not found: {xlsx_path}")
    if not json_path.exists():
        sys.exit(f"JSON file not found: {json_path}")

    data = load_json(json_path)
    existing_titles_lower = [d["title"].lower().strip() for d in data]
    next_id = max((d["id"] for d in data), default=0) + 1

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["To Add"]
    done_ws = wb["Already Added"] if "Already Added" in wb.sheetnames else wb.create_sheet("Already Added")

    # Find the header row by looking for the "title" cell rather than
    # assuming a fixed row number -- the sheet has been reformatted at
    # least once already (the original "HOW TO USE" instructions row and
    # EXAMPLE row were removed), and hardcoding row 2 silently skipped the
    # new row 2's hymn instead of erroring, which is worse than just
    # searching for the header.
    header_row_idx = None
    for r in range(1, min(5, ws.max_row) + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().lower() == "title":
            header_row_idx = r
            break
    if header_row_idx is None:
        sys.exit('Could not find the header row (a row with "title" in column A) in the first 5 rows.')
    first_data_row = header_row_idx + 1

    ready_rows = []      # (row_index, record) to merge
    pending_count = 0
    skipped_examples = 0

    for r in range(first_data_row, ws.max_row + 1):
        values = tuple(ws.cell(row=r, column=c).value for c in range(1, len(HEADERS) + 1))
        if all(v in (None, "") for v in values):
            continue
        record = row_to_dict(values)
        if is_example_row(record):
            skipped_examples += 1
            continue
        if is_ready(record):
            ready_rows.append((r, record))
        else:
            pending_count += 1

    if not ready_rows:
        print(f"No rows are ready to merge yet ({pending_count} pending, need at least title + verses).")
        return

    new_records = []
    warnings = []
    for row_idx, record in ready_rows:
        title = record["title"]
        dupes = find_near_duplicates(title, existing_titles_lower)
        if dupes:
            warnings.append(f'  "{title}" looks similar to existing: {dupes} - check before assuming it is new.')

        hymn = {
            "id": next_id,
            "title": title,
            "tune": record.get("tune") or "",
            "verses": split_verses(record.get("verses") or ""),
            "Author": record.get("Author") or "",
            "Composer": record.get("Composer") or "",
            "midi_url": record.get("midi_url") or "",
            "mp3_url": record.get("mp3_url") or "",
            "credits": record.get("credits") or "",
        }
        new_records.append((row_idx, hymn))
        existing_titles_lower.append(title.lower().strip())
        next_id += 1

    print(f"Ready to merge {len(new_records)} hymn(s):")
    for _, h in new_records:
        print(f'  [{h["id"]}] {h["title"]}')
    if warnings:
        print("\nPossible duplicates (not blocked, just flagged):")
        print("\n".join(warnings))
    if pending_count:
        print(f"\n{pending_count} row(s) still pending (missing title or verses) - left untouched.")

    if args.dry_run:
        print("\n--dry-run: nothing was written.")
        return

    # merge + sort
    all_hymns = data + [h for _, h in new_records]
    all_hymns.sort(key=lambda d: d["title"].lower())
    json_path.write_text(json.dumps(all_hymns, indent=2, ensure_ascii=False), encoding="utf-8")

    # log to "Already Added" and remove merged rows from "To Add" (bottom-up so indices stay valid)
    log_row = done_ws.max_row + 1
    today = date.today().isoformat()
    for _, h in sorted(new_records, key=lambda t: t[0]):
        done_ws.cell(row=log_row, column=1, value=h["title"])
        done_ws.cell(row=log_row, column=2, value=h["id"])
        done_ws.cell(row=log_row, column=3, value=today)
        log_row += 1

    for row_idx, _ in sorted(new_records, key=lambda t: t[0], reverse=True):
        ws.delete_rows(row_idx, 1)

    wb.save(xlsx_path)

    print(f"\nDone. {json_path} now has {len(all_hymns)} hymns, sorted alphabetically.")
    print(f"{len(new_records)} row(s) moved from 'To Add' to 'Already Added' in {xlsx_path}.")


if __name__ == "__main__":
    main()
